import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
import io

st.set_page_config(page_title="SWIMS Order & Inventory Summary", layout="wide")

st.title("📊 SWIMS Order & Inventory Summary Generator")
st.write("Upload your 3-tab Excel file below (Tab 1: Orders | Tab 2: Inventory | Tab 3: Production) to generate the consolidated Summary tab with live Excel formulas.")

uploaded_file = st.file_uploader("Upload Excel Workbook", type=["xlsx", "xls"])

if uploaded_file is not None:
    if st.button("🚀 Generate Summary Tab", type="primary"):
        try:
            with st.spinner("Processing workbook and generating formulas..."):
                file_bytes = uploaded_file.getvalue()
                xls = pd.ExcelFile(io.BytesIO(file_bytes))

                # Read tabs by index position (0, 1, 2)
                df_orders = pd.read_excel(xls, sheet_name=0)
                df_inventory = pd.read_excel(xls, sheet_name=1)
                df_production = pd.read_excel(xls, sheet_name=2)

                # Clean header whitespace
                df_orders.columns = df_orders.columns.str.strip()
                df_inventory.columns = df_inventory.columns.str.strip()
                df_production.columns = df_production.columns.str.strip()

                UPC_COL = 'Upc No'
                for df in [df_orders, df_inventory, df_production]:
                    if UPC_COL in df.columns:
                        df[UPC_COL] = df[UPC_COL].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)

                # 1. Process Orders Pivot
                orders_grouped = df_orders.groupby(UPC_COL).agg({
                    'Style': 'first',
                    'Clr': 'first',
                    'Size Name': 'first',
                    'Open Qty': 'sum',
                    'Open Value': 'sum'
                }).reset_index()

                orders_grouped.rename(columns={
                    'Open Qty': 'Sum of Open Qty',
                    'Open Value': 'Sum of Open Value'
                }, inplace=True)

                # Filter for Open Qty > 0 to calculate active Item Price
                active_orders = df_orders[df_orders['Open Qty'] > 0]
                if not active_orders.empty:
                    price_grouped = active_orders.groupby(UPC_COL)['Price'].mean().reset_index()
                    price_grouped.rename(columns={'Price': 'Item Price'}, inplace=True)
                    orders_grouped = pd.merge(orders_grouped, price_grouped, on=UPC_COL, how='left')
                else:
                    orders_grouped['Item Price'] = 0

                if 'Item Price' not in orders_grouped.columns:
                    fallback_price = df_orders.groupby(UPC_COL)['Price'].mean().reset_index()
                    fallback_price.rename(columns={'Price': 'Item Price'}, inplace=True)
                    orders_grouped = pd.merge(orders_grouped, fallback_price, on=UPC_COL, how='left')

                orders_grouped['Item Price'] = orders_grouped['Item Price'].fillna(0)

                # 2. Process On-Hand Inventory
                if 'On Hand' in df_inventory.columns:
                    inv_grouped = df_inventory.groupby(UPC_COL)['On Hand'].sum().reset_index()
                    inv_grouped.rename(columns={'On Hand': 'On Hand Units'}, inplace=True)
                else:
                    inv_grouped = pd.DataFrame(columns=[UPC_COL, 'On Hand Units'])

                # 3. Process WIP Production
                if 'Po Qty' in df_production.columns:
                    prod_grouped = df_production.groupby(UPC_COL)['Po Qty'].sum().reset_index()
                    prod_grouped.rename(columns={'Po Qty': 'In WIP'}, inplace=True)
                else:
                    prod_grouped = pd.DataFrame(columns=[UPC_COL, 'In WIP'])

                # Merge datasets
                summary_df = pd.merge(orders_grouped, inv_grouped, on=UPC_COL, how='left')
                summary_df = pd.merge(summary_df, prod_grouped, on=UPC_COL, how='left')

                summary_df['On Hand Units'] = summary_df['On Hand Units'].fillna(0).astype(int)
                summary_df['In WIP'] = summary_df['In WIP'].fillna(0).astype(int)

                # 4. Open original workbook and insert new Summary tab
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
                if 'Summary' in wb.sheetnames:
                    del wb['Summary']

                ws = wb.create_sheet(title="Summary", index=0)

                # KPI Top Headers
                ws['A1'] = "Can be fulfilled till end of October"
                ws['A2'] = "Cannot be fulfilled till end of October"

                last_row = len(summary_df) + 6
                ws['E5'] = f"=SUM(E7:E{last_row})"
                ws['F5'] = f"=SUM(F7:F{last_row})"
                ws['G5'] = f"=SUM(G7:G{last_row})"
                ws['H5'] = f"=SUM(H7:H{last_row})"
                ws['I5'] = f"=SUM(I7:I{last_row})"
                ws['K5'] = f"=SUM(K7:K{last_row})"
                ws['L5'] = f"=SUM(L7:L{last_row})"

                ws['C1'] = f"=SUMIF(I7:I{last_row}, \">=0\", E7:E{last_row})"
                ws['D1'] = f"=SUMIF(I7:I{last_row}, \">=0\", F7:F{last_row})"
                ws['C2'] = f"=SUMIF(I7:I{last_row}, \"<0\", E7:E{last_row})"
                ws['D2'] = f"=SUMIF(I7:I{last_row}, \"<0\", F7:F{last_row})"

                headers = [
                    'Upc No', 'Style', 'Clr', 'Size Name',
                    'Sum of Open Qty', 'Sum of Open Value',
                    'On Hand Units', 'In WIP', 'True Availability',
                    'Item Price', 'Null Value', 'Fillable Value'
                ]
                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=6, column=col_idx, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="D9D9D9", fill_type="solid")

                # Data rows and dynamic formulas
                for idx, row in summary_df.iterrows():
                    r = idx + 7
                    ws.cell(row=r, column=1, value=str(row['Upc No']))
                    ws.cell(row=r, column=2, value=row['Style'])
                    ws.cell(row=r, column=3, value=row['Clr'])
                    ws.cell(row=r, column=4, value=row['Size Name'])
                    ws.cell(row=r, column=5, value=row['Sum of Open Qty'])
                    ws.cell(row=r, column=6, value=row['Sum of Open Value'])
                    ws.cell(row=r, column=7, value=row['On Hand Units'])
                    ws.cell(row=r, column=8, value=row['In WIP'])
                    
                    ws.cell(row=r, column=9, value=f"=(G{r}+H{r})-E{r}")
                    ws.cell(row=r, column=10, value=f"=IFERROR(F{r}/E{r}, 0)")
                    ws.cell(row=r, column=11, value=f"=IF(I{r}<0, I{r}*J{r}, 0)")
                    ws.cell(row=r, column=12, value=f"=IF(I{r}>=0, F{r}, 0)")

                # Save output stream
                output = io.BytesIO()
                wb.save(output)
                processed_data = output.getvalue()

            st.success("✅ Summary tab successfully generated with live Excel formulas!")

            st.download_button(
                label="📥 Download Updated Excel Workbook",
                data=processed_data,
                file_name=f"Updated_{uploaded_file.name}",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"An error occurred while processing: {e}")